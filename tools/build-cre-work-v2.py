"""Build original public diagnostic packets; private references are never read."""
import copy, hashlib, importlib.util, json, math, subprocess
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.platypus import Image

ROOT=Path(__file__).resolve().parents[1]; BASE=ROOT/'benchmarks/cre-work-v2'
spec=importlib.util.spec_from_file_location('author_pdf',ROOT/'tools/build-workflow-cases.py'); pdfmod=importlib.util.module_from_spec(spec);spec.loader.exec_module(pdfmod)
p,pdf,table=pdfmod.p,pdfmod.pdf,pdfmod.table

def save(path,data):
 path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(data,indent=2,allow_nan=False)+'\n')
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def case(cid,name,asset,city,units,area,rent,vacant,pending,flags,rate,cap,ltv,payoff):
 return dict(id=cid,name=name,asset=asset,city=city,units=units,area=area,rent=rent,vacant=vacant,pending=pending,flags=flags,rate=rate,cap=cap,ltv=ltv,payoff=payoff)

CASES=[
case('mf-01','Magnolia Terrace','Multifamily','Raleigh, NC',48,820,1650,3,1,['amendment'],.0625,.0575,.65,4200000),
case('of-01','Beacon Plaza','Office','Charlotte, NC',12,5500,11200,2,1,['rollover'],.0675,.075,.60,6500000),
case('rt-01','Maple Crossing','Retail','Columbus, OH',10,3200,6800,1,0,['amendment','reimbursements'],.065,.0675,.65,4300000),
case('in-01','Falcon Distribution','Industrial','Indianapolis, IN',8,22000,14600,1,0,['floor'],.061,.0625,.65,7800000),
case('mf-02','Willow Gardens','Multifamily','San Antonio, TX',64,940,1725,11,5,['leaseup'],.069,.06,.60,7200000),
case('of-02','Summit Exchange','Office','Denver, CO',9,7100,13800,1,0,['free_rent','amendment'],.0725,.08,.60,7400000),
case('rt-02','Crescent Shops','Retail','Tampa, FL',14,2400,6100,2,1,['reimbursements','missing_insurance'],.0675,.07,.60,5500000),
case('in-02','Copper Yard','Industrial','Dallas, TX',6,31500,22400,0,0,['rollover'],.065,.06,.65,9500000),
case('mf-03','Brookside Residences','Multifamily','Richmond, VA',40,780,1530,2,0,['capex'],.0615,.06,.65,3200000),
case('of-03','Harbor Offices','Office','Baltimore, MD',15,4300,8800,4,1,['leaseup','missing_insurance'],.072,.0825,.55,7000000),
case('rt-03','Elm Street Center','Retail','Louisville, KY',8,4200,7500,0,0,['rollover','unsigned'],.067,.0725,.65,3800000),
case('in-03','Pioneer Flex Park','Industrial','Columbus, OH',16,8200,7100,2,0,['capex','reimbursements'],.064,.065,.65,6400000),
case('mf-04','Spruce Court','Multifamily','Jacksonville, FL',72,1010,1850,4,2,['missing_insurance','scan'],.066,.0575,.65,8300000),
case('of-04','Union Professional','Office','Minneapolis, MN',11,6100,12100,1,1,['unsigned','scan'],.069,.0775,.60,6800000),
case('rt-04','Linden Commons','Retail','Phoenix, AZ',12,3700,9200,1,0,['free_rent','floor'],.062,.0675,.65,6600000),
case('in-04','Westgate Logistics','Industrial','Memphis, TN',7,27500,16700,0,0,['scan','floor'],.058,.0675,.65,6900000),
case('mf-05','Palisade Apartments','Multifamily','Greenville, SC',36,870,1580,0,0,['clean'],.06,.06,.60,1800000),
case('of-05','Hillside Business Park','Office','Madison, WI',8,4600,9100,0,0,['clean'],.064,.075,.60,2300000),
case('rt-05','Parkway Village','Retail','Kansas City, MO',9,3100,5900,0,0,['clean'],.0625,.07,.60,1700000),
case('in-05','Cobalt Commerce','Industrial','Louisville, KY',5,21000,14200,0,0,['clean'],.062,.065,.60,3100000),
]
RISK_IDS=['pending_in_occupancy','near_term_rollover','missing_insurance','capital_in_opex','unsigned_rent_change','temporary_free_rent','sizing_floor_binding','refinance_shortfall']


def rows_for(c,revised=False):
 rows=[]
 for i in range(c['units']):
  group=i%3; area=c['area']+[-120,0,210][group];rent=c['rent']+[-150,0,275][group]
  status='vacant' if i>=c['units']-c['vacant'] else 'pending' if i>=c['units']-c['vacant']-c['pending'] else 'active'
  expiry='2027-02-28' if 'rollover' in c['flags'] and i==0 else '2029-08-31'
  if revised and i==(c['units']-c['vacant']-1 if c['pending'] else 1):status='active';rent+=225
  rows.append(dict(unit=f'{i+1:03d}',type=(['Studio','1BR','2BR'][group] if c['asset']=='Multifamily' else ['Small suite','Standard suite','Large suite'][group]),area=area,status=status,rent=rent,expiry=expiry))
 return rows


def truth(c,revised=False):
 rows=rows_for(c,revised);active=[r for r in rows if r['status']=='active'];rents={r['unit']:r['rent'] for r in active}
 if 'amendment' in c['flags']:rents['001']+=350
 annual=sum(rents.values())*12; gross=sum(r['rent'] for r in rows)*12
 egi=gross*.95+c['units']*350
 base_gross=sum(r['rent'] for r in rows_for(c,False))*12
 tax=round(base_gross*.105);ins=round(base_gross*.035);other=round(base_gross*.14)
 rate=max(c['rate'],.07 if 'floor' in c['flags'] else .06)+(.0075 if revised else 0)
 cap=c['cap']+(.0025 if revised else 0);mgmt=egi*.03
 noi=egi-tax-ins-other-mgmt;ncf=noi-c['units']*300
 constant=(rate/12)/(1-(1+rate/12)**(-360))*12
 value=noi/cap;ltv=value*c['ltv'];dscr=ncf/(1.25*constant);dy=ncf/.09;loan=min(ltv,dscr,dy);cash=loan*.99-c['payoff']-45000
 fields={'total_units':len(rows),'occupied_units':len(active),'pending_units':sum(r['status']=='pending' for r in rows),'total_area_sf':sum(r['area'] for r in rows),'occupied_area_sf':sum(r['area'] for r in active),'unit_occupancy':len(active)/len(rows),'area_occupancy':sum(r['area'] for r in active)/sum(r['area'] for r in rows),'annual_in_place_rent':annual,'unit_001_monthly_rent':rents.get('001'),'gross_potential_rent':gross,'uw_egi':egi,'uw_tax':tax,'uw_insurance':ins,'uw_operating_expenses':tax+ins+other+mgmt,'uw_noi':noi,'uw_ncf':ncf,'sizing_rate':rate,'cap_rate':cap,'capitalization_value':value,'ltv_limit':ltv,'dscr_limit':dscr,'debt_yield_limit':dy,'maximum_loan':loan,'net_cash_out':cash,'insurance_policy_limit':None if 'missing_insurance' in c['flags'] else 2000000,'selected_roll_date':'2026-09-01' if revised else '2026-08-31','reported_t12_noi':round(base_gross*.9)-tax-ins-other-(40000 if 'capex' in c['flags'] else 0)}
 for i,typ in enumerate(['type_a','type_b','type_c']):
  group=[r for r in rows if (int(r['unit'])-1)%3==i];fields[typ+'_units']=len(group);fields[typ+'_occupied']=sum(r['status']=='active' for r in group)
 risks={'pending_in_occupancy':bool(c['pending']),'near_term_rollover':'rollover' in c['flags'],'missing_insurance':'missing_insurance' in c['flags'],'capital_in_opex':'capex' in c['flags'],'unsigned_rent_change':'unsigned' in c['flags'],'temporary_free_rent':'free_rent' in c['flags'],'sizing_floor_binding':'floor' in c['flags'],'refinance_shortfall':cash<0}
 for key,val in risks.items():fields['risk_'+key]='present' if val else 'absent'
 return fields,rows


def workbook(path,rows,date):
 w=openpyxl.Workbook();s=w.active;s.title='Rent Roll';s.merge_cells('A1:G1');s['A1']='Property manager certified rent roll';s['A2']='As of';s['B2']=date
 s.append([]);s.append(['Suite / Unit','Unit type','Area SF','Possession status','Monthly base rent','Lease expiration','Note'])
 for r in rows:s.append([r['unit'],r['type'],r['area'],r['status'],r['rent'],r['expiry'],'Pending means not yet in possession' if r['status']=='pending' else ''])
 s.freeze_panes='C5';s.auto_filter.ref=f'A4:G{s.max_row}';s.sheet_view.showGridLines=False
 for col,width in zip('ABCDEFG',[16,22,14,22,23,23,48]):s.column_dimensions[col].width=width
 for cell in s[4]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='183B35')
 for row in s.iter_rows(min_row=5):row[4].number_format='$#,##0.00';row[6].alignment=Alignment(wrap_text=True)
 w.save(path)


def build(c):
 d=BASE/'cases'/c['id'];sources=d/'sources';revision=d/'revision';sources.mkdir(parents=True,exist_ok=True);revision.mkdir(exist_ok=True)
 f,rows=truth(c);rf,rr=truth(c,True);gross=f['gross_potential_rent'];flags=c['flags']
 workbook(sources/'rent-roll-certified.xlsx',rows,'2026-08-31')
 old=copy.deepcopy(rows)
 for r in old:r['rent']-=75
 workbook(sources/'rent-roll-archive.xlsx',old,'2026-06-30')
 tax=f['uw_tax'];ins=f['uw_insurance'];other=round(gross*.14);capex=40000 if 'capex' in flags else 0
 ledger=[['Account','12-month actual USD'],['Rental and ancillary collections',round(gross*.9)],['Real estate tax expense',tax],['Insurance expense',ins],['Other recurring operating expenses',other],['Roof replacement booked in repairs',capex],['Reported NOI',f['reported_t12_noi']]]
 pdf(sources/'operating-statement.pdf',c['name']+' | trailing 12 months',[[p('Period: 2025-09-01 through 2026-08-31. Cash operating statement; parentheses denote expenses. Amounts in USD, not thousands.'),table(ledger,[350,175]),p('The roof replacement is a capital improvement; no recurring repair component is included. Historical collections are not an annualized rent-roll total.')]])
 rent=rows[0]['rent'];lease=[p('Executed lease. Tenant: Original Fixture Tenant 001 LLC. Suite 001.'),p(f'Monthly contractual base rent: ${rent:,.0f}. Expiration: {rows[0]["expiry"]}. Security deposit: ${rent*2:,.0f}. Area: {rows[0]["area"]:,} SF.'),p('A signed later amendment controls conflicting lease terms. An unsigned proposal does not change this agreement.'),p('Base rent is fully abated during September and October 2026; contractual scheduled rent resumes November 1.' if 'free_rent' in flags else 'No free-rent concession applies as of the evaluation date.')]
 pdf(sources/'lease-001.pdf',c['name']+' | Lease excerpt',[[*lease]])
 pdf(sources/'lease-correspondence.pdf',c['name']+' | Correspondence',[[p('Signed amendment effective August 1, 2026. Both parties executed July 20, 2026. Monthly rent increases by $350; all other terms unchanged.' if 'amendment' in flags else 'Unsigned proposal dated August 25, 2026: monthly rent for Suite 001 would increase by $500. No party has signed; no effective change exists.' if 'unsigned' in flags else 'Property manager confirmation dated August 31, 2026: no amendments or new rent concessions affect Suite 001.')]])
 pdf(sources/'insurance.pdf',c['name']+' | Insurance record',[[p('Current insurance expense is supported by the operating statement. The policy declaration page was not supplied; coverage limit and expiry are unavailable. Do not equate annual expense with coverage.' if 'missing_insurance' in flags else 'Policy declaration: general liability aggregate limit $2,000,000. Effective July 1, 2026; expiry June 30, 2027. This coverage limit is distinct from the annual insurance expense in the operating statement.')]])
 pdf(sources/'manager-cover.pdf',c['name']+' | Manager cover note',[[p(f'Property: {c["name"]}. Market: {c["city"]}. Asset: {c["asset"]}. Original fictional evaluation property; no actual address, sponsor credentials or transaction is represented.'),p(f'Manager headline counts occupied plus pending as leased: {f["occupied_units"]+f["pending_units"]} of {c["units"]}. Pending occupants have not taken possession. Certified detail controls physical occupancy. Current contractual rent is subject to executed amendments.'),p('Unit groups A/B/C correspond in order to '+(', '.join(['Studio','1BR','2BR'] if c['asset']=='Multifamily' else ['Small suite','Standard suite','Large suite']))+'.'),p('Approved copy to preserve exactly in the revision: "Source-backed analysis. Assumptions remain subject to diligence."')]])
 pdf(sources/'underwriting-instructions.pdf',c['name']+' | Acquisition analysis instructions',[[p('Prepare an investment/sales OM, not merely a debt-request letter. Use provided facts; photos, market sales/rent comps, sponsor history and asking price are not supplied. Label their absence rather than inventing them.'),p('Underwriting uses all-suite scheduled monthly rent on the current certified roll ×12 (before Suite 001 amendment), a 5% economic loss, and annual other income of $350 per unit/suite. The amendment changes in-place rent only; this prescribed stabilized gross potential is a separate assumption. Reimbursement income is already included; do not add it twice.'),p(f'Use tax ${tax:,.0f}; insurance expense ${ins:,.0f}; other recurring expenses ${other:,.0f}; management 3% of underwritten EGI; replacement reserve $300 per unit/suite. Capitalize NOI before reserves at {c["cap"]:.4%}. Size DSCR and debt yield on NCF after reserves.'),p(f'Loan constraints: maximum LTV {c["ltv"]:.2%}; minimum DSCR 1.25; minimum NCF debt yield 9%; amortization 30 years, monthly payments. Note rate {c["rate"]:.4%}; underwriting rate floor {(.07 if "floor" in flags else .06):.2%}. Size on the greater rate, even when the first 24 months are interest-only. Origination fee 1% of sized loan; closing costs $45,000; existing payoff ${c["payoff"]:,.0f}. Negative net cash-out is additional equity required.'),p('Assess each listed risk independently; a known resolvable source conflict should be explained, not left unresolved. Near-term rollover means an active suite expiry in the next 12 months. No lease expiration beyond that horizon is a near-term issue.')]])
 pdf(sources/'diligence-status.pdf',c['name']+' | Diligence status',[[p('Legal, environmental, property condition and title reviews have not been completed. They are routine pending diligence, not findings of a defect. No evidence of contamination, litigation, fraud or structural failure is supplied. Do not assert such findings.'),p('This packet contains a stale archive and contemporaneous controlling documents. Match documents to this property and their stated dates. Historical NOI, underwritten NOI and NCF have different definitions; label each.')]])
 if 'scan' in flags:
  target=sources/'lease-correspondence.pdf';prefix=d/'temporary-scan'
  subprocess.run(['pdftoppm','-singlefile','-r','120','-png',str(target),str(prefix)],check=True,capture_output=True)
  from reportlab.pdfgen import canvas
  cv=canvas.Canvas(str(target),pagesize=(612,792),invariant=1);cv.drawImage(str(prefix.with_suffix('.png')),0,0,width=612,height=792);cv.save();prefix.with_suffix('.png').unlink()
 workbook(revision/'rent-roll-revised.xlsx',rr,'2026-09-01')
 pdf(revision/'revision-notice.pdf',c['name']+' | Revision instruction',[[p(f"As of September 1, 2026 the attached revised certified roll supersedes the prior roll. Suite {c['units']-c['vacant'] if c['pending'] else 2:03d} is now active and its scheduled monthly rent has increased by $225. All other lease details remain unchanged. Prior executed Suite 001 amendments continue to apply."),p('Increase the INITIAL SIZING RATE by 75 basis points and the capitalization rate by 25 basis points. Use the new roll for all-unit scheduled gross potential. All other underwriting rules and payoff remain unchanged. Update the existing underwriting workbook, OM, reconciled facts and risk table. Preserve the approved copy exactly; identify what changed.')]])
 field_ids=list(f)
 brief=f'''# {c['name']}: prepare a complete investment package\n\nAs of 2026-08-31, analyze the attached original fictional deal packet. Produce a reconciled underwriting XLSX with live formulas, a concise investment/sales offering memorandum PDF, and an evidence-backed final analysis. Both files must be downloadable. Include executive summary, property and unit/suite mix, lease/occupancy analysis, historical versus underwritten financials, debt sizing and sources/uses, risks and missing diligence, and source notes. Do not invent photos, asking price, market comps or sponsor history. Routine generation and outline approvals are preauthorized. Use only these sources; no external research or outreach.\n\nThe workbook must expose input cells for gross_potential_rent, cap_rate, sizing_rate and output cells for uw_egi, uw_noi, uw_ncf, capitalization_value, maximum_loan and net_cash_out. Provide their Sheet!A1 locations. Include a sensitivity table and a source register.\n\nReport the following named fields with value, unit where needed, and evidence filename plus page/cell location. Ratios are decimal values (0.95, not 95). For type_a/b/c use the manager's group mapping. For risk fields use present, absent or unresolved. Missing information remains explicitly unknown, distinct from zero.\n\n{', '.join(field_ids)}\n\nFor risk_pending_in_occupancy assess the manager's inclusion of unpossessed pending units, not whether your final calculation repeats the mistake. For risk_capital_in_opex assess the reported historical ledger classification. For risk_sizing_floor_binding compare the original note rate and floor. For risk_unsigned_rent_change assess whether an unsigned rent proposal is in the packet.\n\nIn your final analysis provide fields as an array of objects with id, value, evidence (array), calculation; qualifications (array); summary; blocked (array). Also supply input_map and output_map. This is an interoperability format; substantive financial content is the assessment.\n'''
 (d/'brief.md').write_text(brief)
 (d/'revision.md').write_text('Continue the same deal and conversation. Apply the attached revision notice and revised rent roll. Deliver revised XLSX and OM PDF and the same named analysis fields and cell maps. Preserve all unaffected facts and the approved sentence. Explain the changes. Do not replace your prior outputs in the audit archive.\n')
 save(d/'reference.json',{'case_id':c['id'],'qualification':'author-reviewed synthetic diagnostic; independent review pending','initial':f,'revision':rf,'critical_fields':['annual_in_place_rent','unit_occupancy','uw_noi','uw_ncf','maximum_loan','net_cash_out'],'risks':RISK_IDS,'expected_unchanged':[k for k in f if f[k]==rf[k]],'approved_copy':'Source-backed analysis. Assumptions remain subject to diligence.'})
 save(d/'authoring.json',c)
 save(d/'manifest.json',{'case_id':c['id'],'asset':c['asset'],'scenario_flags':flags,'source_kind':'original synthetic','independent_deal':False,'shared_generation_family':c['asset'],'as_of':'2026-08-31','revision_as_of':'2026-09-01','initial_files':{x.name:sha(x) for x in sorted(sources.iterdir())},'revision_files':{x.name:sha(x) for x in sorted(revision.iterdir())},'reference_sha256':sha(d/'reference.json'),'brief_sha256':sha(d/'brief.md'),'license':'MIT','review_status':'independent-review-pending'})

if __name__=='__main__':
 for c in CASES:build(c)
 save(BASE/'case-index.json',{'version':'cre-work-v2','qualification':'synthetic diagnostic expansion; not authentic-customer or independently-qualified evidence','cases':[dict(id=c['id'],name=c['name'],asset=c['asset'],flags=c['flags']) for c in CASES],'shared_generation_families':4,'real_deals':0})
 print(json.dumps({'packets':len(CASES),'initial_source_files':sum(len(list((BASE/'cases'/c['id']/'sources').iterdir())) for c in CASES),'revision_source_files':40,'real_deals':0}))
