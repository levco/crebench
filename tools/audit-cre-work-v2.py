"""Read delivered files; all spreadsheet recalculation/perturbation uses copies."""
import hashlib,json,math,re,sys
from pathlib import Path
import openpyxl
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from crebench.work_v2_tools import read_cell,recalc
from crebench.grade_work_v2 import equivalent
from crebench.work_v2_scoring_v1_1 import score as grade
BASE=ROOT/'benchmarks/cre-work-v2';OUT=ROOT/'experiments/2026-09-07-cre-work-v2'

def save(p,x):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(x,indent=2,allow_nan=False)+'\n')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def expectations(c,key,changes):
 gpr=changes.get('gross_potential_rent',key['gross_potential_rent']);egi=gpr*.95+c['units']*350
 other=round(key['gross_potential_rent']*.14)
 # Recurring cost amounts are original fixed inputs, even at the revision stage.
 original=json.loads((BASE/'cases'/c['id']/'reference.json').read_text())['initial'];other=round(original['gross_potential_rent']*.14)
 noi=egi-key['uw_tax']-key['uw_insurance']-other-egi*.03;ncf=noi-c['units']*300
 rate=changes.get('sizing_rate',key['sizing_rate']);cap=changes.get('cap_rate',key['cap_rate']);constant=(rate/12)/(1-(1+rate/12)**-360)*12
 value=noi/cap;loan=min(value*c['ltv'],ncf/(1.25*constant),ncf/.09)
 return {'uw_egi':egi,'uw_noi':noi,'uw_ncf':ncf,'capitalization_value':value,'maximum_loan':loan,'net_cash_out':loan*.99-c['payoff']-45000}

def audit(run,stage):
 result=json.loads((run/f'result-{stage}.json').read_text());case=BASE/'cases'/result['case_id'];ref=json.loads((case/'reference.json').read_text());key=ref[stage];data=json.loads((case/'authoring.json').read_text());analysis_path=run/f'analysis-{stage}.json';analysis=json.loads(analysis_path.read_text()) if analysis_path.exists() else {}
 out={'stage':stage,'workbook':{},'om':{},'professional_acceptance':None,'human_review':'pending','original_hashes':{}}
 wp=run/result['latest_workbook'] if result.get('latest_workbook') else None
 if wp and wp.exists():
  before=sha(wp);out['original_hashes'][str(wp.relative_to(run))]=before;maps=analysis
  spec=wp.with_suffix('.json')
  if spec.exists():maps={**json.loads(spec.read_text()),**{k:v for k,v in analysis.items() if k in ('input_map','output_map') and v}}
  inp=maps.get('input_map',{});outputs=maps.get('output_map',{});folder=run/'audits'/stage
  try:
   book=openpyxl.load_workbook(wp,data_only=False);cached=openpyxl.load_workbook(recalc(wp,folder/'base'),data_only=True);checks=[]
   for field,expected in expectations(data,key,{}).items():
    try:actual=read_cell(cached,outputs[field]);formula=read_cell(book,outputs[field]);passed=equivalent(field,actual,expected)
    except Exception as exc:actual=None;formula=None;passed=False
    checks.append({'field':field,'actual':actual,'expected':expected,'passed':passed,'live_formula':isinstance(formula,str) and formula.startswith('=')})
   perturb=[]
   for name,change in [('cap_rate',key['cap_rate']+.005),('gross_potential_rent',key['gross_potential_rent']*.95),('sizing_rate',key['sizing_rate']+.01)]:
    try:
     copy=openpyxl.load_workbook(wp,data_only=False);sheet,cell=inp[name].rsplit('!',1);copy[sheet.strip("'")][cell]=change
     changed=folder/f'{name}.xlsx';copy.save(changed);calc=openpyxl.load_workbook(recalc(changed,folder/name),data_only=True);expected=expectations(data,key,{name:change});items=[]
     for field,value in expected.items():
      try:actual=read_cell(calc,outputs[field]);passed=equivalent(field,actual,value)
      except Exception:actual=None;passed=False
      items.append({'field':field,'actual':actual,'expected':value,'passed':passed})
     perturb.append({'input':name,'checks':items,'passed':all(i['passed'] for i in items)})
    except Exception as exc:perturb.append({'input':name,'passed':False,'error':str(exc)})
   out['workbook']={'readable':True,'base_checks':checks,'perturbations':perturb,'all_base_outputs_correct':all(c['passed'] for c in checks),'all_sensitivities_correct':all(p['passed'] for p in perturb)}
  except Exception as exc:out['workbook']={'readable':False,'error':str(exc)}
  assert sha(wp)==before,'Original workbook changed during audit'
 else:out['workbook']={'readable':False,'reason':'No recorded delivered workbook'}
 mp=run/result['latest_memo'] if result.get('latest_memo') else None
 if mp and mp.exists():
  before=sha(mp);out['original_hashes'][str(mp.relative_to(run))]=before
  try:
   reader=PdfReader(mp);text='\n'.join(page.extract_text() or '' for page in reader.pages);(run/f'om-{stage}.txt').write_text(text)
   out['om']={'readable':bool(text.strip()),'pages':len(reader.pages),'characters':len(text),'approved_copy_present':ref['approved_copy'] in re.sub(r'\s+',' ',text),'financial_and_presentation_review':'pending; readable PDF is not acceptance'}
  except Exception as exc:out['om']={'readable':False,'error':str(exc)}
  assert sha(mp)==before,'Original PDF changed during audit'
 else:out['om']={'readable':False,'reason':'No recorded delivered PDF'}
 if stage=='revision':
  g=grade(case,analysis,stage);unchanged=set(ref['expected_unchanged'])
  if ref['initial']['pending_units']>0 and ref['revision']['pending_units']==0:unchanged.discard('risk_pending_in_occupancy')
  checks=[r for r in g['checks'] if r['field'] in unchanged and r.get('scored',True)];out['unchanged_facts']={'passed':sum(c['passed'] for c in checks),'total':len(checks)}
  initial_path=run/'result-initial.json'
  initial=json.loads(initial_path.read_text()) if initial_path.exists() else {}
  out['revision_files_changed']={}
  for label,field in [('workbook','latest_workbook'),('om','latest_memo')]:
   old=run/initial[field] if initial.get(field) else None;new=run/result[field] if result.get(field) else None
   out['revision_files_changed'][label]=bool(old and new and old.exists() and new.exists() and sha(old)!=sha(new))
 save(run/f'artifact-audit-{stage}.json',out);return out

if __name__=='__main__':
 counts={}
 for root in [OUT/'credential-resumption/api',OUT/'native',OUT/'consumer']:
  for p in root.glob('**/result-initial.json'):
   for stage in ('initial','revision'):
    if not (p.parent/f'result-{stage}.json').exists():continue
    audit_file=p.parent/f'artifact-audit-{stage}.json'
    if audit_file.exists():continue
    a=audit(p.parent,stage);counts[stage]=counts.get(stage,0)+1
 print(json.dumps(counts))
