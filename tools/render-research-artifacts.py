"""Render original comp-grid excerpts without rewriting any model workbook."""
import concurrent.futures
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
EXP = ROOT / 'experiments/2026-09-07-research-v1'
OUT = ROOT / 'work/research-v1/artifact-previews'
OUT.mkdir(parents=True, exist_ok=True)
node = sys.argv[1] if len(sys.argv) > 1 else 'node'
jobs = []
for file in EXP.glob('*/*/*/artifact-audit.json'):
    audit = json.loads(file.read_text())
    if not (audit.get('file') or '').endswith('.xlsx') or '-comps-' not in file.parent.name:
        continue
    workbook = file.parent / audit['file']
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    stem = '-'.join(file.parent.relative_to(EXP).parts)
    image = OUT / (stem + '.png')
    evidence = OUT / (stem + '.json')
    if image.exists() and evidence.exists() and json.loads(evidence.read_text()).get('sha256') == digest:
        continue
    book = openpyxl.load_workbook(workbook)
    sheet = book[book.sheetnames[0]]
    region = f'A1:{get_column_letter(min(sheet.max_column, 24))}{min(sheet.max_row, 12)}'
    jobs.append((workbook, sheet.title, region, image, evidence, digest))

def render(job):
    workbook, sheet, region, image, evidence, digest = job
    result = subprocess.run([node, str(ROOT / 'tools/render-workflow-workbook.mjs'), str(workbook), sheet, region, str(image)],
                            cwd=ROOT, capture_output=True, text=True, timeout=120)
    row = dict(workbook=str(workbook.relative_to(ROOT)), sha256=digest, sheet=sheet, range=region,
               image=str(image.relative_to(ROOT)), rendered=result.returncode == 0,
               note='Original workbook excerpt only; rendering does not rewrite or repair the file.')
    if result.returncode:
        row['error'] = result.stderr[-1000:]
    evidence.write_text(json.dumps(row, indent=2) + '\n')
    print(json.dumps(row), flush=True)

with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
    list(pool.map(render, jobs))
