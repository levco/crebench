"""Independently recalculate originals and published perturbation copies.

Only public Python dependencies and LibreOffice are needed. No model calls and
no artifact-tool dependency. Source workbooks and formulas are never rewritten.
Outputs must go to a new local directory; the published audit is not overwritten.
"""
import argparse,concurrent.futures,hashlib,importlib.util,json,math,subprocess,sys
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
spec=importlib.util.spec_from_file_location('audit',ROOT/'tools/verify-workflow-artifact.py');audit=importlib.util.module_from_spec(spec);spec.loader.exec_module(audit)
from crebench.grade_workflow import equivalent


def replay(run,out):
    r=json.loads((run/'result.json').read_text());original=run/r['latest_workbook'];sha=hashlib.sha256(original.read_bytes()).hexdigest()
    prior=json.loads((run/'audit/artifact-audit.json').read_text())
    if prior.get('source_workbook_sha256')!=sha:raise ValueError('Original workbook differs from audited SHA-256: '+str(run))
    mapping=run/'location-map.json' if (run/'location-map.json').exists() else original.with_suffix('.json')
    maps=json.loads(mapping.read_text());outputs=maps['output_map'];case=ROOT/'benchmarks/workflow-v1/cases'/r['case_id'];c=json.loads((case/'authoring-data.json').read_text());key=json.loads((case/'answer-key.json').read_text())['fields'];checks=[]
    original_formulas=openpyxl.load_workbook(original,data_only=False)
    base=openpyxl.load_workbook(audit.recalc(original,out/'base'),data_only=True)
    for field in ['uw_noi_before_reserves','uw_ncf','maximum_loan']:
        actual=audit.read_cell(base,outputs[field]);checks.append({'scenario':'base','field':field,'actual':actual,'expected':key[field],'passed':equivalent(field,actual,key[field])})
    for scenario in ['cap','rate','egi']:
        archived=run/'audit'/f'{scenario}.xlsx';changes=json.loads((run/'audit'/f'{scenario}-changes.json').read_text());copy=openpyxl.load_workbook(archived,data_only=False)
        # Recalculate first: some exported numeric edits retain an empty <f/> node.
        # LibreOffice resolves that representation to the supplied numeric value.
        recalculated=openpyxl.load_workbook(audit.recalc(archived,out/scenario),data_only=True)
        # Verify the effective edits and preservation of all remaining formulas.
        for location,value in changes.items():
            actual=audit.read_cell(recalculated,location)
            if not isinstance(actual,(int,float)) or not math.isclose(actual,value,rel_tol=1e-12,abs_tol=1e-10):raise ValueError('Perturbation edit mismatch: '+location)
        for sheet in original_formulas:
            for row in sheet:
                for cell in row:
                    location=f'{sheet.title}!{cell.coordinate}'
                    if cell.data_type=='f' and location not in changes:
                        actual=copy[sheet.title][cell.coordinate].value
                        canonical=lambda s:str(s).replace('_xlfn.','').replace('_xlws.','').replace('$','').replace(' ','').upper()
                        if canonical(actual)!=canonical(cell.value):raise ValueError('Unexpected formula change: '+location)
        expected=audit.expected_scenario(c,key,scenario)
        fields={'cap':['capitalization_value','ltv_limit','maximum_loan'],'rate':['dscr_limit','maximum_loan'],'egi':['uw_noi_before_reserves','uw_ncf','maximum_loan']}[scenario]
        for field in fields:
            actual=audit.read_cell(recalculated,outputs[field]);checks.append({'scenario':scenario,'field':field,'actual':actual,'expected':expected[field],'passed':equivalent(field,actual,expected[field]),'edited_cells':changes})
    errors=[{'cell':f'{s.title}!{cell.coordinate}','error':cell.value} for s in base for row in s for cell in row if cell.data_type=='e']
    result={'run':str(run.relative_to(ROOT)),'original_sha256':sha,'numeric_checks':checks,'formula_errors':errors,'passed_numeric_checks':sum(x['passed'] for x in checks),'total_numeric_checks':len(checks),'source_formula_preservation_verified':True}
    (out/'replay.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps({k:result[k] for k in ['run','passed_numeric_checks','total_numeric_checks']}),flush=True);return result


def main():
    p=argparse.ArgumentParser();p.add_argument('experiment');p.add_argument('--output',required=True);a=p.parse_args();exp=Path(a.experiment).resolve();out=Path(a.output).resolve();out.mkdir(parents=True,exist_ok=False)
    version=subprocess.run(['soffice','--version'],capture_output=True,text=True,check=True).stdout.strip();jobs=[]
    for record in exp.glob('*/*/*/result.json'):
        r=json.loads(record.read_text())
        if r.get('latest_workbook') and (record.parent/'audit/artifact-audit.json').exists():jobs.append((record.parent,out/record.parent.relative_to(exp)))
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda job:replay(*job),jobs))
    summary={'engine':version,'runs':len(results),'checks':sum(r['total_numeric_checks'] for r in results),'passed':sum(r['passed_numeric_checks'] for r in results),'workbooks_with_formula_errors':sum(bool(r['formula_errors']) for r in results),'results':results};(out/'summary.json').write_text(json.dumps(summary,indent=2)+'\n');print(json.dumps({k:v for k,v in summary.items() if k!='results'}))
    if summary['checks']!=summary['passed']:sys.exit(1)
if __name__=='__main__':main()
