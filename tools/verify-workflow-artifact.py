"""Audit delivered XLSX formulas and perturbations without repairing the output.

All edits are made on copies through artifact-tool. Recalculate with LibreOffice.
Location maps come from the evaluated system (or an attributed reviewer mapping),
never from the gold workbook's layout. Missing maps remain an auditable limitation.
"""
import argparse
import json
import math
from pathlib import Path
import subprocess
import sys
import zipfile
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from crebench.grade_workflow import equivalent


def read_cell(book,location):
    sheet,cell=location.rsplit('!',1);return book[sheet.strip("'")][cell].value


def pv_capacity(c,ncf,rate):
    return sum((ncf/c['dscr']/12)/(1+rate/12)**i for i in range(1,c['amort']*12+1))


def expected_scenario(c,base,scenario):
    egi=base['uw_egi']*(.95 if scenario=='egi' else 1)
    management=max(c['expenses'][-1],egi*c['mgmt'])
    noi=egi-(base['uw_operating_expenses']-base['uw_management_fee']+management)
    ncf=noi-base['uw_reserves'];rate=base['sizing_rate']+(.01 if scenario=='rate' else 0)
    cap=c['cap']+(.005 if scenario=='cap' else 0);value=noi/cap
    ltv=value*c['ltv'];dscr=pv_capacity(c,ncf,rate);dy=(noi if c['dy_basis']=='NOI' else ncf)/c['dy']
    return {'uw_egi':egi,'uw_noi_before_reserves':noi,'uw_ncf':ncf,'capitalization_value':value,
            'ltv_limit':ltv,'dscr_limit':dscr,'debt_yield_limit':dy,'maximum_loan':min(ltv,dscr,dy)}


def recalc(path,folder):
    folder.mkdir(parents=True,exist_ok=True)
    result=subprocess.run(['soffice','-env:UserInstallation='+ (folder/'lo-profile').resolve().as_uri(),
        '--headless','--convert-to','xlsx','--outdir',str(folder),str(path)],capture_output=True,text=True,timeout=90)
    (folder/'recalc.log').write_text(result.stdout+result.stderr)
    target=folder/path.name
    if result.returncode or not target.exists():raise RuntimeError('LibreOffice did not produce a workbook: '+result.stderr[-500:])
    return target


def verify(case,workbook,mapping,directory,node):
    case=Path(case);workbook=Path(workbook);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    c=json.loads((case/'authoring-data.json').read_text());key=json.loads((case/'answer-key.json').read_text())['fields']
    maps=json.loads(Path(mapping).read_text());inputs=maps.get('input_map',{});outputs=maps.get('output_map',{})
    result={'case_id':case.name,'source_workbook':workbook.name,'mapping_source':Path(mapping).name,
            'engine':'LibreOfficeDev 26.8.0.0.alpha0','checks':[],'scenarios':{},'manual_checks_pending':[2,3,10]}
    def check(id,passed,detail):result['checks'].append({'rubric_check':id,'passed':passed,'detail':detail})
    try:
        formula_book=openpyxl.load_workbook(workbook,data_only=False);check(1,True,'XLSX opens')
        basebook=openpyxl.load_workbook(recalc(workbook,directory/'base'),data_only=True)
    except Exception as exc:
        check(1,False,str(exc));return result
    expected_base=['uw_noi_before_reserves','uw_ncf','maximum_loan']
    base_checks=[]
    for field in expected_base:
        try:value=read_cell(basebook,outputs[field]);passed=equivalent(field,value,key[field])
        except Exception as exc:value=None;passed=False
        base_checks.append({'field':field,'actual':value,'expected':key[field],'passed':passed})
    check(4,all(x['passed'] for x in base_checks),base_checks)
    formula_outputs={}
    for field in ['uw_noi_before_reserves','uw_ncf','maximum_loan','capitalization_value','dscr_limit']:
        try:formula_outputs[field]=isinstance(read_cell(formula_book,outputs[field]),str) and read_cell(formula_book,outputs[field]).startswith('=')
        except Exception:formula_outputs[field]=False
    check(5,all(formula_outputs.values()),formula_outputs)
    # Each independent perturbation starts from the original, not another scenario.
    for scenario,id in [('cap',6),('rate',7),('egi',8)]:
        try:
            if scenario=='cap':
                loc=inputs.get('cap_rate');old=read_cell(basebook,loc);changes={loc:old+.005}
            elif scenario=='rate':
                loc=inputs.get('sizing_rate_increment') or inputs.get('sizing_rate');old=read_cell(basebook,loc);changes={loc:old+.01}
            else:
                if inputs.get('egi_gross_potential_rent') and inputs.get('egi_other_income'):
                    locations=[inputs['egi_gross_potential_rent'],inputs['egi_other_income']]
                    changes={loc:read_cell(basebook,loc)*.95 for loc in locations}
                else:
                    loc=inputs.get('egi_multiplier') or inputs.get('egi');old=read_cell(basebook,loc);changes={loc:old*.95}
            expected=expected_scenario(c,key,scenario);config=directory/f'{scenario}-changes.json';config.write_text(json.dumps(changes,indent=2)+'\n')
            modified=directory/f'{scenario}.xlsx'
            build=subprocess.run([node,str(ROOT/'tools/workflow-perturb.mjs'),str(workbook),str(config),str(modified)],capture_output=True,text=True,timeout=90)
            (directory/f'{scenario}-patch.log').write_text(build.stdout+build.stderr)
            if build.returncode:raise RuntimeError(build.stderr[-800:])
            w=openpyxl.load_workbook(recalc(modified,directory/scenario),data_only=True)
            fields={'cap':['capitalization_value','ltv_limit','maximum_loan'],
                    'rate':['dscr_limit','maximum_loan'],
                    'egi':['uw_noi_before_reserves','uw_ncf','maximum_loan']}[scenario]
            tests=[]
            for field in fields:
                actual=read_cell(w,outputs[field]);tests.append({'field':field,'actual':actual,'expected':expected[field],'passed':equivalent(field,actual,expected[field])})
            check(id,all(t['passed'] for t in tests),tests);result['scenarios'][scenario]={'changes':changes,'checks':tests}
        except Exception as exc:
            check(id,False,str(exc));result['scenarios'][scenario]={'error':str(exc),'classification':'Unverified or failed perturbation; inspect detail before interpreting as a financial error'}
    errors=[f'{s.title}!{cell.coordinate}: {cell.value}' for s in basebook for row in s for cell in row if cell.data_type=='e']
    with zipfile.ZipFile(workbook) as z:
        external=[name for name in z.namelist() if name.startswith('xl/externalLinks/') or name.endswith('vbaProject.bin')]
    check(9,not errors and not external,{'formula_errors':errors,'external_links_or_macros':external})
    result['automatic_checks_passed']=sum(x['passed'] for x in result['checks']);result['automatic_checks_total']=len(result['checks'])
    return result


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('case');p.add_argument('workbook');p.add_argument('mapping');p.add_argument('directory');p.add_argument('--node',default='node');a=p.parse_args()
    result=verify(a.case,a.workbook,a.mapping,a.directory,a.node);(Path(a.directory)/'artifact-audit.json').write_text(json.dumps(result,indent=2,allow_nan=False)+'\n')
    print(json.dumps({'case':result['case_id'],'checks':[(x['rubric_check'],x['passed']) for x in result['checks']],'manual_checks_pending':result['manual_checks_pending']}))
