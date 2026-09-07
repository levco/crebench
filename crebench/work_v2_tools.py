"""Public, bounded document tools. No proprietary artifact package is required."""
import copy,json,re,subprocess
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils.cell import coordinate_to_tuple
from .workflow_tools import ToolContext,TOOLS as BASE_TOOLS,source_text

TOOLS=copy.deepcopy(BASE_TOOLS)
for t in TOOLS:
 f=t['function']
 if f['name']=='create_memo':f['description']=f['description'].replace('financing memorandum','offering memorandum or other requested document')
 if f['name']=='submit_analysis':
  f['parameters']['properties'].update(input_map={'type':'object','additionalProperties':{'type':'string'}},output_map={'type':'object','additionalProperties':{'type':'string'}})


def read_cell(book,location):
 sheet,cell=location.rsplit('!',1);return book[sheet.strip("'")][cell].value

def recalc(path,folder):
 folder.mkdir(parents=True,exist_ok=True)
 result=subprocess.run(['soffice','-env:UserInstallation='+(folder/'profile').resolve().as_uri(),'--headless','--convert-to','xlsx','--outdir',str(folder),str(path)],capture_output=True,text=True,timeout=90)
 (folder/'recalc.log').write_text(result.stdout+result.stderr)
 target=folder/path.name
 if result.returncode or not target.exists():raise ValueError('Spreadsheet recalculation failed: '+result.stderr[-500:])
 return target


def write_workbook(spec,path):
 sheets=spec.get('sheets',[])
 if not 1<=len(sheets)<=12:raise ValueError('1 to 12 sheets required')
 if sum(len(s.get('cells',[])) for s in sheets)>6000:raise ValueError('6000 populated-cell limit')
 book=openpyxl.Workbook();book.remove(book.active)
 for sheet in sheets:
  name=sheet['name']
  if len(name)>31 or re.search(r'[\[\]:*?/\\]',name) or name in book.sheetnames:raise ValueError('Unique valid worksheet names required')
  book.create_sheet(name)
 for sheet in sheets:
  s=book[sheet['name']];s.sheet_view.showGridLines=False
  for col,width in sheet.get('column_widths',{}).items():
   if not re.fullmatch('[A-Z]{1,3}',col):raise ValueError('Column letters required')
   s.column_dimensions[col].width=min(90,max(8,float(width)))
  seen=set()
  for item in sheet['cells']:
   addr=item['cell'];row,col=coordinate_to_tuple(addr)
   if row>5000 or col>100 or addr in seen:raise ValueError('Unique cells within 5000 rows and 100 columns required')
   seen.add(addr);cell=s[addr]
   if 'formula' in item:
    formula=item['formula']
    if not formula.startswith('=') or re.search(r'\[|\]|https?://|\\',formula,re.I):raise ValueError('Local formulas only')
    allowed={'SUM','MIN','MAX','IF','PMT','PV','ROUND','COUNT','COUNTA','AVERAGE','SUMIF','SUMIFS','COUNTIF','COUNTIFS','IFERROR','ABS'}
    funcs=set(re.findall(r'([A-Za-z_][A-Za-z_0-9.]*)\s*\(',formula))
    if not {f.upper() for f in funcs}<=allowed:raise ValueError('Unsupported function; use scalar arithmetic or '+','.join(sorted(allowed)))
    cell.value=formula
   else:
    cell.value='Unknown' if item.get('value') is None else item['value']
    if isinstance(cell.value,str):cell.data_type='s'
   role=item.get('role','body');cell.font=Font(name='Arial',size=15 if role=='title' else 10,bold=role in ('title','header','total'))
   if role=='header':cell.fill=PatternFill('solid',fgColor='E5EEE9')
   if role=='input':cell.font=Font(name='Arial',size=10,color='185CB5')
   cell.alignment=Alignment(vertical='top',wrap_text=item.get('wrap',True))
   if 'number_format' in item:cell.number_format=item['number_format']
   s.row_dimensions[row].height=max(s.row_dimensions[row].height or 18,min(180,item.get('height',30)))
  if sheet.get('freeze_rows'):s.freeze_panes='A'+str(sheet['freeze_rows']+1)
  s.print_options.horizontalCentered=True;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
 book.save(path)


class WorkContext(ToolContext):
 def __init__(self,case,run,node='node'):
  super().__init__(case,run,node);self.stage='initial';self.latest_cached=None
 def source(self,name):
  if not isinstance(name,str) or Path(name).name!=name:raise ValueError('Source filename only')
  dirs=[self.case/'sources']+([self.case/'revision'] if self.stage=='revision' else [])
  for d in dirs:
   p=d/name
   if p.is_file() and p.suffix in ('.pdf','.xlsx'):return p
  raise ValueError('File not available in this stage')
 def inspect(self):
  if not self.latest_workbook:return {'error':'No workbook created'}
  spec=json.loads(self.latest_spec.read_text());book=openpyxl.load_workbook(self.latest_cached or self.latest_workbook,data_only=True)
  outputs={}
  for key,loc in spec.get('output_map',{}).items():
   try:outputs[key]=read_cell(book,loc)
   except Exception as exc:outputs[key]={'error':str(exc)}
  errors=[{'sheet':s.title,'cell':c.coordinate,'value':c.value} for s in book for row in s for c in row if c.data_type=='e']
  return {'sheets':book.sheetnames,'output_values':outputs,'formula_errors':errors,'file':str(self.latest_workbook.relative_to(self.run)),'recalculated':self.latest_cached is not None}
 def call(self,name,args):
  if name=='list_files':
   dirs=[self.case/'sources']+([self.case/'revision'] if self.stage=='revision' else [])
   return {'files':[{'name':p.name,'bytes':p.stat().st_size} for d in dirs for p in sorted(d.iterdir()) if p.suffix in ('.pdf','.xlsx')]}
  if name=='create_workbook':
   self.versions['workbook']+=1;stem='underwriting-v'+str(self.versions['workbook']);spec=self.run/'artifacts'/f'{stem}.json';output=spec.with_suffix('.xlsx')
   spec.write_text(json.dumps(args['spec'],indent=2)+'\n');write_workbook(args['spec'],output)
   self.latest_spec=spec;self.latest_workbook=output;self.latest_cached=recalc(output,self.run/'recalculated'/stem)
   return {'created':str(output.relative_to(self.run)),'bytes':output.stat().st_size,'inspection':self.inspect()}
  if name=='submit_analysis':
   (self.run/f'analysis-{self.stage}.json').write_text(json.dumps(args,indent=2,allow_nan=False)+'\n');self.done=True
   return {'submitted':True,'stage':self.stage}
  return super().call(name,args)
