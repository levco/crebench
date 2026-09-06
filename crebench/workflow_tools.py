"""Provider-neutral bounded tools for the public workflow agent.

No answer keys, CRE-specific calculators, shell, network, or arbitrary execution
are exposed to evaluated models. Workbook formulas and document text are data.
"""
import ast
import base64
import json
import math
import operator
from pathlib import Path
import re
import subprocess
import sys
from xml.sax.saxutils import escape

import openpyxl
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
CELL_SCHEMA={'type':'object','properties':{
    'cell':{'type':'string','description':'A1 address'},'value':{'type':['string','number','boolean','null']},
    'formula':{'type':'string','description':'Excel formula beginning =; omit value when using formula'},
    'number_format':{'type':'string'},'role':{'type':'string','enum':['body','title','header','input','total','note']},
    'wrap':{'type':'boolean'},'height':{'type':'number'}},'required':['cell']}
SPEC_SCHEMA={'type':'object','properties':{
    'sheets':{'type':'array','items':{'type':'object','properties':{
        'name':{'type':'string'},'cells':{'type':'array','items':CELL_SCHEMA},
        'column_widths':{'type':'object','additionalProperties':{'type':'number'}},'freeze_rows':{'type':'integer'}},'required':['name','cells']}},
    'input_map':{'type':'object','additionalProperties':{'type':'string'},'description':'Identify cap_rate, sizing_rate (or sizing_rate_increment), egi (or egi_multiplier) as Sheet!A1 addresses. Alternatively provide egi_gross_potential_rent and egi_other_income, which can both be scaled with loss percentage held constant.'},
    'output_map':{'type':'object','additionalProperties':{'type':'string'},'description':'Map requested field IDs, especially maximum_loan, uw_noi_before_reserves, uw_ncf, capitalization_value, ltv_limit, dscr_limit, to Sheet!A1 addresses'}},'required':['sheets','input_map','output_map']}


def function(name,description,properties,required=()):
    return {'type':'function','function':{'name':name,'description':description,'parameters':{'type':'object','properties':properties,'required':list(required),'additionalProperties':False}}}


TOOLS=[
function('list_files','List the provided source files. Only the authorized case packet is accessible.',{}),
function('read_file','Read a provided PDF as page-labeled text or an XLSX as coordinate-labeled cell values/formulas. Image-only PDF pages are attached visually for you. No answer key is available.',{'name':{'type':'string'}},['name']),
function('view_pdf_page','Render a page of a provided PDF for visual inspection. The page image is attached as a user image after the tool result.',{'name':{'type':'string'},'page':{'type':'integer','minimum':1}},['name','page']),
function('calculate','Evaluate arithmetic. Supports + - * / **, parentheses, min, max, abs, round, sqrt, log, exp and pmt(rate_per_period, periods, present_value). No variables, code or CRE-specific formulas.',{'expression':{'type':'string'}},['expression']),
function('calculate_batch','Preferred for related calculations: evaluate named arithmetic steps in order. Later expressions can refer to names of earlier results. Same arithmetic/functions as calculate. No file access or domain-specific answers. Example steps [{name:"revenue",expression:"1200*12"},{name:"net",expression:"revenue-4000"}].',{'steps':{'type':'array','items':{'type':'object','properties':{'name':{'type':'string'},'expression':{'type':'string'}},'required':['name','expression'],'additionalProperties':False}}},['steps']),
function('create_workbook','Create a real XLSX from your own cell values and live Excel formulas. Neutral Arial styling is applied according to roles; you control sheet layout, labels, formulas, widths and number formats. All sheets are created before formulas. Use simple scalar Excel formulas (SUM, MIN, MAX, IF, PMT) and quote sheet names with spaces. At most 12 sheets and 6000 populated cells. No macros, external links or web functions. Supply input/output cell maps. You may revise by calling again; all versions remain in the trace.',{'spec':SPEC_SCHEMA},['spec']),
function('inspect_workbook','Inspect the latest created workbook, including cached values, formula errors, sheet names and mapped output values. This tool does not compare to an answer key.',{}),
function('create_memo','Create a real financing memorandum PDF from your supplied prose and tables. Neutral professional styling and automatic pagination are supplied. You control section order, content, headings and tables. Plain text only; do not supply executable HTML.',{
    'title':{'type':'string'},'subtitle':{'type':'string'},'sections':{'type':'array','items':{'type':'object','properties':{
        'heading':{'type':'string'},'paragraphs':{'type':'array','items':{'type':'string'}},
        'table':{'type':'object','properties':{'headers':{'type':'array','items':{'type':'string'}},'rows':{'type':'array','items':{'type':'array','items':{'type':['string','number','null']}}}},'required':['headers','rows']},
        'page_break_before':{'type':'boolean'}},'required':['heading','paragraphs']}}},['title','subtitle','sections']),
function('submit_analysis','Submit your final analysis and requested fields. Include sources and location citations. This records your answer and completes the run; first create requested files and inspect them. Formatting itself is not financially graded.',{
    'fields':{'type':'array','items':{'type':'object','properties':{
        'id':{'type':'string'},'value':{'type':['number','string','null']},'evidence':{'type':'array','items':{'type':'string'}},'calculation':{'type':'string'}},'required':['id','value','evidence','calculation']}},
    'qualifications':{'type':'array','items':{'type':'string'}},'summary':{'type':'string'},'blocked':{'type':'array','items':{'type':'string'}}},['fields','qualifications','summary','blocked'])]


def arithmetic(expression,variables=None):
    if not isinstance(expression,str) or len(expression)>2000:raise ValueError('Expression too long')
    binary={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv,ast.Pow:operator.pow}
    def pmt(rate,n,pv):
        if n<=0:raise ValueError('periods must be positive')
        return -pv/n if rate==0 else -pv*rate/(1-(1+rate)**(-n))
    functions={'min':min,'max':max,'abs':abs,'round':round,'sqrt':math.sqrt,'log':math.log,'exp':math.exp,'pmt':pmt}
    def read(node,depth=0):
        if depth>30:raise ValueError('Expression nesting limit')
        if isinstance(node,ast.Constant) and type(node.value) in (int,float):value=node.value
        elif isinstance(node,ast.Name) and variables is not None and node.id in variables:value=variables[node.id]
        elif isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):value=read(node.operand,depth+1)*(1 if isinstance(node.op,ast.UAdd) else -1)
        elif isinstance(node,ast.BinOp) and type(node.op) in binary:
            left,right=read(node.left,depth+1),read(node.right,depth+1)
            if isinstance(node.op,ast.Pow) and abs(right)>10000:raise ValueError('Exponent limit')
            value=binary[type(node.op)](left,right)
        elif isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id in functions and not node.keywords:
            value=functions[node.func.id](*(read(a,depth+1) for a in node.args))
        else:raise ValueError('Arithmetic only')
        if not isinstance(value,(int,float)) or not math.isfinite(value) or abs(value)>1e18:raise ValueError('Finite result below 1e18 required')
        return value
    return read(ast.parse(expression,mode='eval').body)


def source_text(path):
    if path.suffix=='.pdf':
        return '\n\n'.join(f'PAGE {i+1}\n'+(page.extract_text(extraction_mode='layout') or '[No extractable text; inspect the page image]') for i,page in enumerate(PdfReader(path).pages))
    if path.suffix=='.xlsx':
        w=openpyxl.load_workbook(path,data_only=False,read_only=True);parts=[]
        for s in w:
            parts.append('WORKSHEET '+s.title)
            for row in s:
                values=[f'{c.coordinate}={json.dumps(c.value,default=str)}' for c in row if c.value is not None]
                if values:parts.append(' | '.join(values))
        return '\n'.join(parts)
    raise ValueError('Only source PDF and XLSX files are exposed')


def parse_json(text):
    text=text.strip()
    blocks=re.findall(r'```(?:json)?\s*([\s\S]*?)```',text,flags=re.I)
    if len(blocks)==1:text=blocks[0].strip()
    elif len(blocks)>1:raise ValueError('Multiple JSON blocks require attributed transcription')
    def pairs(items):
        obj={}
        for key,value in items:
            if key in obj:raise ValueError('Duplicate JSON key: '+key)
            obj[key]=value
        return obj
    def constant(value):raise ValueError('Nonfinite JSON value: '+value)
    return json.loads(text,object_pairs_hook=pairs,parse_constant=constant)


class ToolContext:
    def __init__(self,case,run,node='node'):
        self.case=Path(case);self.run=Path(run);self.node=node;self.versions={'workbook':0,'memo':0};self.images=[]
        self.latest_spec=None;self.latest_workbook=None;self.done=False
        (self.run/'artifacts').mkdir(parents=True,exist_ok=True)
        (self.run/'input-views').mkdir(parents=True,exist_ok=True)

    def source(self,name):
        if not isinstance(name,str) or Path(name).name!=name:raise ValueError('Source filename only')
        p=self.case/'sources'/name
        if p.suffix not in ('.pdf','.xlsx') or not p.is_file():raise ValueError('Not an available source file')
        return p

    def page(self,name,page):
        p=self.source(name)
        if p.suffix!='.pdf' or type(page)!=int or not 1<=page<=len(PdfReader(p).pages):raise ValueError('Invalid PDF page')
        prefix=self.run/'input-views'/f'{p.stem}-p{page}'
        subprocess.run(['pdftoppm','-f',str(page),'-singlefile','-scale-to','1800','-png',str(p),str(prefix)],check=True,capture_output=True,timeout=30)
        image=prefix.with_suffix('.png');self.images.append(image)
        return {'file':name,'page':page,'image_path':str(image.relative_to(self.run)),'image_attached':True}

    def call(self,name,args):
        if name=='list_files':return {'files':[{'name':p.name,'bytes':p.stat().st_size} for p in sorted((self.case/'sources').iterdir()) if p.suffix in ('.pdf','.xlsx')]}
        if name=='read_file':
            source=self.source(args['name']);result={'name':source.name,'content':source_text(source),'images':[]}
            if source.suffix=='.pdf':
                for i,page in enumerate(PdfReader(source).pages):
                    if len((page.extract_text() or '').strip())<30:result['images'].append(self.page(source.name,i+1))
            return result
        if name=='view_pdf_page':return self.page(args['name'],args['page'])
        if name=='calculate':return {'expression':args['expression'],'value':arithmetic(args['expression'])}
        if name=='calculate_batch':
            if len(args['steps'])>120:raise ValueError('120-step limit')
            values={}
            for step in args['steps']:
                key=step['name']
                if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,63}',key) or key in values:raise ValueError('Unique simple calculation names required')
                values[key]=arithmetic(step['expression'],values)
            return {'values':values}
        if name=='create_workbook':
            self.versions['workbook']+=1;stem=f'underwriting-v{self.versions["workbook"]}';spec=self.run/'artifacts'/f'{stem}.json';output=spec.with_suffix('.xlsx')
            spec.write_text(json.dumps(args['spec'],indent=2,allow_nan=False)+'\n')
            result=subprocess.run([self.node,str(ROOT/'tools/workflow-workbook.mjs'),str(spec),str(output)],capture_output=True,text=True,timeout=90)
            (self.run/'artifacts'/f'{stem}.build.txt').write_text(result.stdout+result.stderr)
            if result.returncode or not output.exists():raise ValueError('Workbook build failed: '+(result.stderr+result.stdout)[-2500:])
            self.latest_spec=spec;self.latest_workbook=output
            return {'created':str(output.relative_to(self.run)),'bytes':output.stat().st_size,'inspection':self.inspect()}
        if name=='inspect_workbook':return self.inspect()
        if name=='create_memo':
            self.versions['memo']+=1;stem=f'memorandum-v{self.versions["memo"]}';spec=self.run/'artifacts'/f'{stem}.json';output=spec.with_suffix('.pdf')
            spec.write_text(json.dumps(args,indent=2,allow_nan=False)+'\n');write_memo(args,output)
            return {'created':str(output.relative_to(self.run)),'bytes':output.stat().st_size,'pages':len(PdfReader(output).pages),'extracted_text':source_text(output)}
        if name=='submit_analysis':
            ids=[f['id'] for f in args['fields']]
            if len(ids)!=len(set(ids)):raise ValueError('Duplicate field IDs; resolve before submitting')
            (self.run/'analysis.json').write_text(json.dumps(args,indent=2,allow_nan=False)+'\n');self.done=True
            return {'recorded':True,'fields':len(ids),'workbooks':self.versions['workbook'],'memos':self.versions['memo']}
        raise ValueError('Unknown tool')

    def inspect(self):
        if not self.latest_workbook:raise ValueError('No workbook created yet')
        w=openpyxl.load_workbook(self.latest_workbook,data_only=True);f=openpyxl.load_workbook(self.latest_workbook,data_only=False)
        errors=[f'{s.title}!{c.coordinate}: {c.value}' for s in w for row in s for c in row if c.data_type=='e']
        mappings=json.loads(self.latest_spec.read_text()).get('output_map',{})
        outputs={}
        for name,loc in mappings.items():
            try:
                sheet,cell=loc.rsplit('!',1);outputs[name]=w[sheet.strip("'")][cell].value
            except Exception as exc:outputs[name]={'error':str(exc)}
        return {'sheets':w.sheetnames,'formula_count':sum(c.data_type=='f' for s in f for row in s for c in row),'formula_errors':errors[:40],'outputs':outputs}


def write_memo(spec,path):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pairs=[('/System/Library/Fonts/Supplemental/Arial.ttf','/System/Library/Fonts/Supplemental/Arial Bold.ttf'),('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')]
    regular,bold=next(pair for pair in pairs if all(Path(p).exists() for p in pair))
    pdfmetrics.registerFont(TTFont('MemoBody',regular));pdfmetrics.registerFont(TTFont('MemoBold',bold))
    styles=getSampleStyleSheet()
    for s in styles.byName.values():s.fontName='MemoBold' if s.name in ['Title','Heading1','Heading2'] else 'MemoBody'
    styles.add(ParagraphStyle(name='MemoCell',fontName='MemoBody',fontSize=9,leading=12))
    styles['BodyText'].fontSize=10;styles['BodyText'].leading=15;styles['BodyText'].spaceAfter=10
    para=lambda t,s='BodyText':Paragraph(escape(str(t)).replace('\n','<br/>'),styles[s])
    story=[para(spec['title'],'Title'),para(spec['subtitle']),Spacer(1,12)]
    if len(json.dumps(spec))>100000 or len(spec['sections'])>24:raise ValueError('Memo length limit')
    for section in spec['sections']:
        if section.get('page_break_before'):story.append(PageBreak())
        story.append(para(section['heading'],'Heading2'))
        story.extend(para(t) for t in section['paragraphs'])
        if section.get('table'):
            data=section['table'];n=len(data['headers'])
            if not 1<=n<=6 or any(len(row)!=n for row in data['rows']):raise ValueError('Table must have 1-6 consistent columns')
            rows=[[para(x,'MemoCell') for x in row] for row in [data['headers']]+data['rows']]
            t=Table(rows,colWidths=[504/n]*n,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EEE9')),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#7E9387')),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.extend([t,Spacer(1,12)])
    def footer(c,doc):
        c.setFont('MemoBody',8);c.setFillColor(colors.HexColor('#536057'));c.drawString(54,27,'Financing memorandum | Internal review | Fictional evaluation case');c.drawRightString(558,27,str(doc.page))
    SimpleDocTemplate(str(path),pagesize=(612,792),leftMargin=54,rightMargin=54,topMargin=44,bottomMargin=44,title=spec['title'],author='CRE Bench evaluated system').build(story,onFirstPage=footer,onLaterPages=footer)
